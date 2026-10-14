"""BOM spreadsheet parser — port of feature/search/BomSpreadsheetParser.kt.

Reads the first worksheet of a ``.xlsx`` BOM file and extracts rows keyed by
the canonical LCSC BOM headers (No., Quantity, Comment, Designator, Footprint,
Value, Manufacturer Part, Manufacturer, Supplier Part, Supplier).

The Android version uses Apache POI's ``WorkbookFactory`` (which also reads
legacy ``.xls``). openpyxl only supports ``.xlsx``; legacy ``.xls`` files must
be re-saved as ``.xlsx`` first — an acceptable desktop trade-off.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from ..domain.bom import BomSearchEntry, ParsedBomDocument

_EXPECTED_HEADERS = [
    "No.", "Quantity", "Comment", "Designator", "Footprint", "Value",
    "Manufacturer Part", "Manufacturer", "Supplier Part", "Supplier",
]


def _cell_as_string(cell) -> str | None:
    if cell is None or cell.value is None:
        return None
    value = cell.value
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)):
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value)
    text = str(value).strip()
    return text or None


def parse_bom_file(file_path: str | Path) -> ParsedBomDocument:
    """Parse a BOM spreadsheet. Raises ValueError with a human message on failure."""
    path = Path(file_path)
    workbook = load_workbook(filename=str(path), data_only=True)
    try:
        if not workbook.sheetnames:
            raise ValueError("The BOM file does not contain any worksheet.")
        sheet = workbook[workbook.sheetnames[0]]
        header_row = sheet[1] if sheet.max_row >= 1 else None
        if header_row is None:
            raise ValueError("The BOM file is missing a header row.")

        header_index: dict[str, int] = {}
        for cell in header_row:
            name = _cell_as_string(cell)
            if name:
                header_index[name] = cell.column

        entries: list[BomSearchEntry] = []
        for row_index in range(2, sheet.max_row + 1):
            row = sheet[row_index]
            entry = _row_to_bom_entry(row, header_index, row_index)
            if entry is not None:
                entries.append(entry)

        if not entries:
            raise ValueError("The BOM file does not contain any recognizable data rows.")

        return ParsedBomDocument(file_name=path.name, entries=entries)
    finally:
        workbook.close()


def _row_to_bom_entry(row, header_index: dict[str, int], row_index: int) -> BomSearchEntry | None:
    def get(header: str) -> str | None:
        col = header_index.get(header)
        if col is None:
            return None
        cell = row[col - 1] if col - 1 < len(row) else None
        return _cell_as_string(cell)

    def get_int(header: str) -> int | None:
        raw = get(header)
        if raw is None:
            return None
        try:
            return int(float(raw))
        except (ValueError, TypeError):
            return None

    row_number = get("No.")
    quantity = get_int("Quantity")
    comment = get("Comment")
    designator = get("Designator")
    footprint = get("Footprint")
    value = get("Value")
    manufacturer_part = get("Manufacturer Part")
    manufacturer = get("Manufacturer")
    supplier_part = get("Supplier Part")
    supplier = get("Supplier")

    if (
        (row_number is None or not row_number.strip())
        and quantity is None
        and (comment is None or not comment.strip())
        and (designator is None or not designator.strip())
        and (footprint is None or not footprint.strip())
        and (value is None or not value.strip())
        and (manufacturer_part is None or not manufacturer_part.strip())
        and (manufacturer is None or not manufacturer.strip())
        and (supplier_part is None or not supplier_part.strip())
        and (supplier is None or not supplier.strip())
    ):
        return None

    return BomSearchEntry(
        row_number=row_number or str(row_index),
        quantity=quantity,
        comment=comment,
        designator=designator,
        footprint=footprint,
        value=value,
        manufacturer_part=manufacturer_part,
        manufacturer=manufacturer,
        supplier_part=supplier_part,
        supplier=supplier,
    )
