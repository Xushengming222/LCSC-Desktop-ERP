"""BOM result list exporter — exports the current BOM search results to .xlsx.

Exports all rows (or the currently filtered subset: All / Matched / Unmatched)
as an Excel workbook with a single sheet.  The sheet mirrors the canonical
LCSC BOM headers so that the exported file can be re-imported as a BOM later,
plus extra columns for match status and matched part numbers.
"""
from __future__ import annotations

import time
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from ..domain.bom import BomMatchFilter, BomSearchRowUiModel

# Canonical LCSC BOM headers (re-import compatible), followed by export extras.
_HEADERS = [
    "No.",
    "Quantity",
    "Comment",
    "Designator",
    "Footprint",
    "Value",
    "Manufacturer Part",
    "Manufacturer",
    "Supplier Part",
    "Supplier",
    "Match Status",
    "Matched Part Number",
    "Matched MPN",
    "Matched Name",
    "Matched Brand",
    "Matched Package",
    "Matched Location(s)",
]

_STATUS_ALL = "ALL"
_STATUS_MATCHED = "Matched"
_STATUS_UNMATCHED = "Unmatched"


def _match_status_text(row: BomSearchRowUiModel) -> str:
    if row.is_bound or bool(row.matched_results):
        return _STATUS_MATCHED
    return _STATUS_UNMATCHED


def _filtered_rows(
    rows: list[BomSearchRowUiModel],
    bom_filter: BomMatchFilter,
) -> list[BomSearchRowUiModel]:
    if bom_filter == BomMatchFilter.MATCHED:
        return [r for r in rows if (r.is_bound or bool(r.matched_results))]
    if bom_filter == BomMatchFilter.UNMATCHED:
        return [r for r in rows if not (r.is_bound or bool(r.matched_results))]
    return list(rows)


def build_bom_export_file_name(bom_filter: BomMatchFilter, source_file_name: str = "") -> str:
    timestamp = time.strftime("%Y%m%d_%H%M%S")
    filter_tag = {
        BomMatchFilter.ALL: "all",
        BomMatchFilter.MATCHED: "matched",
        BomMatchFilter.UNMATCHED: "unmatched",
    }.get(bom_filter, "all")
    base = Path(source_file_name).stem if source_file_name else "bom"
    return f"{base}_{filter_tag}_{timestamp}.xlsx"


def export_bom_rows_to_path(
    rows: list[BomSearchRowUiModel],
    bom_filter: BomMatchFilter,
    path: str | Path,
    source_file_name: str = "",
) -> str | None:
    """Export the given BOM rows (filtered) to an .xlsx file.

    Returns ``None`` on success, or an error message on failure.
    """
    try:
        filtered = _filtered_rows(rows, bom_filter)

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = {
            BomMatchFilter.ALL: "BOM-ALL",
            BomMatchFilter.MATCHED: "BOM-Matched",
            BomMatchFilter.UNMATCHED: "BOM-Unmatched",
        }.get(bom_filter, "BOM")

        # Header row
        for col_index, header in enumerate(_HEADERS, start=1):
            sheet.cell(row=1, column=col_index, value=header)

        # Data rows
        for row_index, row in enumerate(filtered, start=2):
            entry = row.entry
            # Primary match: first matched result
            primary = row.matched_results[0] if row.matched_results else None

            # Matched locations summary: "Code(Qty); Code(Qty)"
            if primary and primary.locations:
                loc_summary = "; ".join(
                    f"{loc.code}({loc.quantity})" for loc in primary.locations
                )
            else:
                loc_summary = ""

            values = [
                entry.row_number,
                entry.quantity if entry.quantity is not None else "",
                entry.comment or "",
                entry.designator or "",
                entry.footprint or "",
                entry.value or "",
                entry.manufacturer_part or "",
                entry.manufacturer or "",
                entry.supplier_part or "",
                entry.supplier or "",
                _match_status_text(row),
                primary.part_number if primary else "",
                primary.mpn or "" if primary else "",
                primary.name or "" if primary else "",
                primary.brand or "" if primary else "",
                primary.package_name or "" if primary else "",
                loc_summary,
            ]
            for col_index, value in enumerate(values, start=1):
                sheet.cell(row=row_index, column=col_index, value=value)

        # Column width auto-fit (rough but readable)
        for col_index in range(1, len(_HEADERS) + 1):
            column_letter = get_column_letter(col_index)
            max_len = len(str(_HEADERS[col_index - 1]))
            for row_index in range(2, len(filtered) + 2):
                cell = sheet.cell(row=row_index, column=col_index)
                cell_value = cell.value
                if cell_value is None:
                    continue
                length = len(str(cell_value))
                if length > max_len:
                    max_len = length
            # Cap width so columns don't grow absurdly wide
            width = min(max(12, max_len + 2), 48)
            sheet.column_dimensions[column_letter].width = width

        # Freeze the header row for easier scrolling
        sheet.freeze_panes = "A2"

        workbook.save(str(path))
        return None
    except Exception as error:  # noqa: BLE001
        return str(error) or "BOM export failed"
